"""DOCX text extraction service."""
from typing import Dict, List, Any
from io import BytesIO

from docx import Document


def extract_text_from_docx(file_content: bytes) -> Dict[str, Any]:
    """
    Extract text from DOCX file.
    
    Args:
        file_content: Raw bytes of the DOCX file
        
    Returns:
        Dictionary with extracted text and metadata
    """
    try:
        # Load document from bytes
        doc = Document(BytesIO(file_content))
        
        # Extract text from paragraphs
        paragraphs = []
        full_text_parts = []
        
        for i, para in enumerate(doc.paragraphs):
            text = para.text.strip()
            if text:  # Only include non-empty paragraphs
                paragraphs.append({
                    "index": i,
                    "text": text,
                })
                full_text_parts.append(text)
        
        # Join all text
        full_text = "\n\n".join(full_text_parts)
        
        # Extract text from tables
        tables_text = []
        for table_idx, table in enumerate(doc.tables):
            table_data = []
            for row in table.rows:
                row_data = [cell.text.strip() for cell in row.cells]
                table_data.append(" | ".join(row_data))
            
            if table_data:
                tables_text.append(f"\n[Table {table_idx + 1}]\n" + "\n".join(table_data))
        
        # Add tables to full text if any
        if tables_text:
            full_text += "\n\n" + "\n".join(tables_text)
        
        # Basic metadata
        metadata = {
            "paragraph_count": len(paragraphs),
            "table_count": len(doc.tables),
            "section_count": len(doc.sections),
        }
        
        # Check if document is empty
        if not full_text or full_text.strip() == "":
            raise ValueError(
                "No text content found in DOCX file. "
                "The document appears to be empty or contains only images/shapes."
            )
        
        return {
            "text": full_text,
            "paragraphs": paragraphs,
            "paragraph_count": len(paragraphs),
            "metadata": metadata,
        }
        
    except Exception as e:
        raise ValueError(f"Failed to extract text from DOCX: {str(e)}")
